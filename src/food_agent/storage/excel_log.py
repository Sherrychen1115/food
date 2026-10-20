"""寫 Excel:一個一直往下加的檔。被占用時先暫存到 pending/,每分鐘重試補寫。"""

import json
import os
import threading
import time
from dataclasses import asdict, dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook

HEADERS = ["日期時間", "上傳者", "食物敘述", "成份", "重量(公克)", "熱量(大卡)", "當日已攝取熱量(大卡)", "照片存放位置"]
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")


@dataclass
class LogRow:
    when: str  # "YYYY-MM-DD HH:MM:SS"(當初處理時的值,補寫也不變)
    uploader: str
    name: str
    ingredients: str
    weight_g: float | None
    calories_kcal: float | None
    daily_total: float
    photo: str  # 相對於 DATA_DIR

    def values(self) -> list:
        return [self.when, self.uploader, self.name, self.ingredients,
                self.weight_g, self.calories_kcal, self.daily_total, self.photo]


def _set_text(cell, text: str) -> None:
    cell.value = text
    # ==== 防護 4：Excel 公式防護　開始 ====
    # 開頭是 = + - @(或 Tab、換行)的文字,強制存成純文字,Excel 打開時不會當公式執行。
    if text.startswith(FORMULA_PREFIXES):
        cell.data_type = "s"
        cell.quotePrefix = True
    # ==== 防護 4 結束 ====


def _append_row(ws, row: LogRow) -> None:
    r = ws.max_row + 1
    for c, value in enumerate(row.values(), start=1):
        cell = ws.cell(row=r, column=c)
        if isinstance(value, str):
            _set_text(cell, value)
        else:
            cell.value = value


class ExcelLog:
    def __init__(self, records_dir: Path, pending_dir: Path):
        self.path = records_dir / "food_log.xlsx"
        self._lock_marker = records_dir / f"~${self.path.name}"  # Excel 開著檔案時會產生
        self._pending_dir = pending_dir
        self._lock = threading.Lock()

    def add(self, row: LogRow) -> bool:
        """先存 pending 再補寫,所以當機也不會遺失。回傳 True 表示已經在 Excel 裡。"""
        with self._lock:
            name = f"{time.time_ns()}.json"
            tmp = self._pending_dir / (name + ".tmp")
            tmp.write_text(json.dumps(asdict(row), ensure_ascii=False), encoding="utf-8")
            os.replace(tmp, self._pending_dir / name)
            return self._flush()

    def flush_pending(self) -> bool:
        with self._lock:
            return self._flush()

    def _pending_files(self) -> list[Path]:
        return sorted(self._pending_dir.glob("*.json"))

    def _load_pending(self) -> list[tuple[Path, LogRow]]:
        loaded = []
        for f in self._pending_files():
            try:
                loaded.append((f, LogRow(**json.loads(f.read_text(encoding="utf-8")))))
            except (ValueError, TypeError, OSError):
                f.rename(f.with_suffix(".bad"))
                print(f"⚠ 暫存檔格式錯誤,已改名為 {f.with_suffix('.bad').name}")
        return loaded

    def _flush(self) -> bool:
        pending = self._load_pending()
        if not pending:
            return True
        if self._lock_marker.exists():
            print(f"⚠ Excel 檔案正被開啟,無法寫入。{len(pending)} 筆已暫存,稍後自動重試。請關閉 {self.path.name}。")
            return False
        try:
            if self.path.exists():
                wb = load_workbook(self.path)
            else:
                wb = Workbook()
                wb.active.append(HEADERS)
            ws = wb.active
            for _, row in pending:
                _append_row(ws, row)
            tmp = self.path.with_suffix(".xlsx.tmp")
            wb.save(tmp)
            os.replace(tmp, self.path)  # 先存暫時檔再取代,寫到一半當機也不會弄壞原檔
        except Exception as e:  # 被占用、權限、檔案壞掉:警告後等下次重試,不能讓程式當掉
            print(f"⚠ 寫入 Excel 失敗({type(e).__name__}),{len(pending)} 筆已暫存,稍後自動重試。")
            return False
        for f, _ in pending:
            f.unlink(missing_ok=True)
        return True

    def sum_for_date(self, date_str: str) -> float:
        """某一天(YYYY-MM-DD)Excel 裡加上尚未補寫的熱量總和。"""
        total = 0.0
        try:
            if self.path.exists():
                wb = load_workbook(self.path, read_only=True)
                for r in wb.active.iter_rows(min_row=2, values_only=True):
                    if str(r[0] or "").startswith(date_str) and isinstance(r[5], (int, float)):
                        total += r[5]
                wb.close()
        except Exception as e:
            print(f"⚠ 讀取 Excel 失敗({type(e).__name__}),今日累計可能偏低。")
        for _, row in self._load_pending():
            if row.when.startswith(date_str) and row.calories_kcal:
                total += row.calories_kcal
        return total
